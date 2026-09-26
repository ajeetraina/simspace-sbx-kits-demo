#!/usr/bin/env python3
"""Export the deck to a 16:9 PowerPoint — one full-bleed slide image per page,
with the speaker notes from gen_deck.py attached to each slide.

    pip install python-pptx
    python3 build_pptx.py   ->  authority-as-code.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches
import gen_deck

HERE = os.path.dirname(os.path.abspath(__file__))
PNG = os.path.join(HERE, "assets", "svg")  # crisp PNGs rendered by build.py

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

for i, (alt, note) in enumerate(gen_deck.SLIDES, 1):
    img = os.path.join(PNG, f"slide-{i:02d}.png")
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(img, 0, 0, width=prs.slide_width, height=prs.slide_height)
    if note:
        slide.notes_slide.notes_text_frame.text = note

out = os.path.join(HERE, "authority-as-code.pptx")
prs.save(out)
print(f"{len(gen_deck.SLIDES)} slides -> {out}")
